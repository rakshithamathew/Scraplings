from hashlib import sha256

from docx import Document
import pytest

from job_automation.database import JobRepository, JobStatus, OutreachStatus, initialize_database
from job_automation.outreach.contact_repository import ContactRepository
from job_automation.outreach.discovery import PublicContact
from job_automation.outreach.personalized import PersonalizedOutreachService, words
from job_automation.resume.resume_parser import parse_resume


@pytest.fixture
def context(tmp_path):
    engine = initialize_database(f"sqlite:///{(tmp_path / 'outreach.db').as_posix()}")
    repo = JobRepository(engine)
    path = tmp_path / 'cv.docx'
    doc = Document()
    for line in ['Alex Candidate', 'Professional Summary', 'Technical Lead with 5+ years of experience building web applications.',
        'Technical Skills', 'React, TypeScript, Python, SQL', 'Work Experience',
        'Technical Lead at Example Health',
        'Implemented a medical document platform to digitize 25+ years of patient records.',
        'Developed a React dashboard for clinical teams, improving data review workflows.',
        'Projects', 'Clinical Dashboard: React and TypeScript application for reviewing patient records.',
        'Tech Stack: React, TypeScript']:
        doc.add_paragraph(line)
    doc.save(path)
    repo.set_active_resume(original_filename='cv.docx', path='cv.docx', content_type='application/vnd.openxmlformats-officedocument.wordprocessingml.document',
        sha256=sha256(path.read_bytes()).hexdigest(), parsed_profile={'years_of_experience': 25, 'skills': ['InventedSkill']})
    job = repo.create_job(title='Frontend Developer', company='Health Company', source='test',
        application_url='https://example.com/jobs/1', status=JobStatus.QUALIFIED,
        description='Build React and TypeScript clinical dashboards to review patient records. Python and SQL are useful. Kubernetes expertise is required.',
        description_complete=True)
    ContactRepository(engine).save(job.id, PublicContact(name='Riley Recruiter', title='Talent Acquisition', company=job.company,
        linkedin_url='https://www.linkedin.com/in/riley', source='https://example.com/team', priority=2, evidence='Talent Acquisition at Health Company'))
    yield repo, job, path, PersonalizedOutreachService(repo, project_root=tmp_path)
    engine.dispose()


def test_generation_grounding_lengths_and_idempotence(context):
    repo, job, path, service = context
    assert parse_resume(path).years_of_experience == 5
    first = service.run()
    assert first.generated == 1
    saved = service.list(job.id)[0]
    assert saved.status == 'DRAFTED'
    assert 150 <= words(saved.email_body) <= 200
    assert 50 <= words(saved.linkedin_message) <= 80
    assert '5+ years of experience' in saved.email_body
    assert '25 years of experience' not in saved.email_body
    assert 'InventedSkill' not in saved.email_body
    assert 'Kubernetes' not in saved.email_body
    assert saved.contact_snapshot['work_email'] is None
    assert 'Riley Recruiter' in saved.email_body
    assert 'Health Company' in saved.email_subject
    assert saved.evidence['cv_excerpt'] in saved.email_body
    assert repo.get_job(job.id).outreach_status is OutreachStatus.DRAFTED
    assert service.run().unchanged == 1
    assert len(service.list()) == 1


def test_missing_jd_or_contact_never_filled_in(context):
    repo, job, _, service = context
    repo.update_job(job.id, description_complete=None)
    summary = service.run()
    assert summary.waiting_input == 1
    row = service.list(job.id)[0]
    assert row.email_body is None
    assert 'Complete job description' in row.missing_inputs[0]
    no_contact = repo.create_job(title=job.title, company=job.company, source='test',
        application_url='https://example.com/jobs/2', status=JobStatus.QUALIFIED,
        description=job.description, description_complete=True)
    service.run(no_contact.id)
    assert service.list(no_contact.id)[0].linkedin_message is None
    assert 'contact' in service.list(no_contact.id)[0].missing_inputs[0]


def test_changed_cv_invalidates_draft_without_using_stale_cache(context):
    repo, job, path, service = context
    service.run()
    doc = Document(path)
    doc.add_paragraph('New content')
    doc.save(path)
    summary = service.run()
    assert summary.waiting_input == 1
    row = service.list(job.id)[0]
    assert row.email_body is None
    assert 'hash' in row.missing_inputs[0]


def test_cv_without_explicit_experience_is_not_assigned_five_years(context):
    repo, job, path, service = context
    doc = Document(path)
    doc.paragraphs[2].text = 'Technical Lead building web applications.'
    doc.save(path)
    repo.set_active_resume(original_filename='cv.docx', path='cv.docx', content_type='docx',
        sha256=sha256(path.read_bytes()).hexdigest(), parsed_profile={})
    service.run()
    assert service.list(job.id)[0].email_body is None
    assert 'years of experience' in service.list(job.id)[0].missing_inputs[0]


def test_only_qualified_jobs_and_outreach_state_preserved(context):
    repo, job, _, service = context
    repo.update_job(job.id, outreach_status=OutreachStatus.SENT)
    service.run()
    assert repo.get_job(job.id).outreach_status is OutreachStatus.SENT
    repo.update_job(job.id, status=JobStatus.APPLIED)
    assert service.run().qualified_jobs == 0
    with pytest.raises(ValueError, match='qualified'):
        service.run(job.id)


def test_jd_change_regenerates_and_delete_cleans_drafts(context):
    repo, job, _, service = context
    service.run()
    old_hash = service.list(job.id)[0].input_hash
    repo.update_job(job.id, description=job.description + ' Additional responsibilities include SQL reporting.', description_complete=True)
    assert service.run().generated == 1
    assert service.list(job.id)[0].input_hash != old_hash
    assert repo.delete_job(job.id)
    assert service.list() == []


def test_outreach_api_returns_stored_drafts(context):
    from fastapi.testclient import TestClient
    from job_automation.api.main import create_app
    repo, job, path, service = context
    app = create_app(str(repo.engine.url), project_root=path.parent)
    with TestClient(app) as client:
        response = client.post(f'/jobs/{job.id}/outreach/generate')
        assert response.status_code == 200
        assert response.json()['generated'] == 1
        draft = client.get(f'/jobs/{job.id}/outreach').json()
        assert draft['status'] == 'DRAFTED'
        assert 150 <= words(draft['email_body']) <= 200
        assert 50 <= words(draft['linkedin_message']) <= 80
        assert client.post('/outreach/generate').json()['unchanged'] == 1
        assert client.post('/jobs/99999/outreach/generate').status_code == 404


def test_changing_description_requires_new_full_detail_evidence(context):
    repo, job, _, service = context
    service.run()
    repo.update_job(job.id, description='React listing snippet')
    assert repo.get_job(job.id).description_complete is False
    service.run()
    assert service.list(job.id)[0].status == 'WAITING_INPUT'
