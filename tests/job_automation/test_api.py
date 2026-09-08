from __future__ import annotations

from datetime import datetime, timezone
from io import BytesIO
from pathlib import Path

import pytest
from fastapi.testclient import TestClient
from docx import Document

from job_automation.api.main import create_app
from job_automation.applications.service import ApplicationCycleSummary
from job_automation.auth import ConnectionState
from job_automation.database import JobRepository, JobStatus, WorkplaceType
from job_automation.matching import UserProfile
from job_automation.scraper.service import DiscoverySummary
from job_automation.resume import ParsedCandidateProfile


@pytest.fixture
def api(tmp_path: Path) -> tuple[TestClient, JobRepository, dict[str, int]]:
    (tmp_path / "resumes").mkdir()
    (tmp_path / "resumes" / "fullstack_resume.pdf").touch()
    app = create_app(f"sqlite:///{(tmp_path / 'api.db').as_posix()}", project_root=tmp_path)
    repository: JobRepository = app.state.repository
    repository.set_active_resume(
        original_filename="fullstack_resume.pdf",
        path="resumes/fullstack_resume.pdf",
        content_type="application/pdf",
        sha256="a" * 64,
        parsed_profile=ParsedCandidateProfile(
            job_titles=["Python Engineer"],
            skills=["Python"],
            technologies=["Python"],
            years_of_experience=3,
            keywords=["backend", "API"],
        ).model_dump(),
    )
    first = repository.create_job(
        title="Python Engineer",
        company="Example",
        location="Remote",
        description="Python backend API role requiring 3 years experience.",
        skills=["Python"],
        source="test",
        application_url="https://example.test/jobs/1",
        match_score=85,
        status=JobStatus.QUALIFIED,
        recommended_resume="resumes/fullstack_resume.pdf",
        workplace_type=WorkplaceType.REMOTE,
        is_open=True,
    )
    second = repository.create_job(
        title="Data Analyst",
        company="Other",
        location="London",
        description="Analyze reports using SQL.",
        skills=["SQL"],
        source="test",
        application_url="https://example.test/jobs/2",
        match_score=55,
        status=JobStatus.DISCOVERED,
        workplace_type=WorkplaceType.ONSITE,
        is_open=True,
    )
    with TestClient(app) as client:
        yield client, repository, {"first": first.id, "second": second.id}


def test_health(api: tuple[TestClient, JobRepository, dict[str, int]]) -> None:
    client, _, _ = api
    response = client.get("/health")
    assert response.status_code == 200
    assert response.json() == {"status": "ok"}


def test_jobs_list_and_filters(api: tuple[TestClient, JobRepository, dict[str, int]]) -> None:
    client, _, _ = api

    assert len(client.get("/jobs").json()) == 2
    qualified = client.get("/jobs", params={"status": "QUALIFIED"})
    high_scores = client.get("/jobs", params={"min_score": 80})

    assert qualified.status_code == 200
    assert [job["title"] for job in qualified.json()] == ["Python Engineer"]
    assert qualified.json()[0]["recommended_resume"] == "resumes/fullstack_resume.pdf"
    assert qualified.json()[0]["outreach_status"] == "NOT_STARTED"
    assert high_scores.status_code == 200
    assert [job["match_score"] for job in high_scores.json()] == [85.0]
    assert client.get("/jobs", params={"status": "INVALID"}).status_code == 422


def test_get_job_and_missing_job(api: tuple[TestClient, JobRepository, dict[str, int]]) -> None:
    client, _, ids = api
    assert client.get(f"/jobs/{ids['first']}").json()["title"] == "Python Engineer"
    assert client.get("/jobs/999999").status_code == 404


def test_patch_status(api: tuple[TestClient, JobRepository, dict[str, int]]) -> None:
    client, repository, ids = api
    response = client.patch(f"/jobs/{ids['second']}/status", json={"status": "SKIPPED"})

    assert response.status_code == 200
    assert response.json()["status"] == "SKIPPED"
    assert repository.get_job(ids["second"]).status is JobStatus.SKIPPED  # type: ignore[union-attr]

    direct_applied = client.patch(f"/jobs/{ids['second']}/status", json={"status": "APPLIED"})
    assert direct_applied.status_code == 400


def test_mark_applied(api: tuple[TestClient, JobRepository, dict[str, int]]) -> None:
    client, repository, ids = api
    applied_at = "2026-09-05T10:30:00Z"
    response = client.post(
        f"/jobs/{ids['first']}/mark-applied",
        json={
            "resume_used": "resumes/fullstack_resume.pdf",
            "application_method": "MANUAL",
            "applied_at": applied_at,
            "application_confirmation": "Confirmation page displayed after submission",
        },
    )

    assert response.status_code == 200
    assert response.json()["status"] == "APPLIED"
    assert response.json()["resume_used"] == "resumes/fullstack_resume.pdf"
    stored = repository.get_job(ids["first"])
    assert stored is not None
    assert stored.status is JobStatus.APPLIED
    assert stored.application_method == "MANUAL"
    assert stored.application_status.value == "APPLIED"
    assert stored.applied_at is not None
    assert stored.application_confirmation is not None

    assert client.post(
        f"/jobs/{ids['first']}/mark-applied",
        json={
            "resume_used": "resumes/fullstack_resume.pdf",
            "application_method": "MANUAL",
            "application_confirmation": "duplicate",
        },
    ).status_code == 409

    invalid = client.post(
        f"/jobs/{ids['second']}/mark-applied",
        json={"resume_used": "resume.pdf", "application_method": "AUTOMATIC"},
    )
    assert invalid.status_code == 422

    blank_resume = client.post(
        f"/jobs/{ids['second']}/mark-applied",
        json={"resume_used": "   ", "application_method": "MANUAL"},
    )
    assert blank_resume.status_code == 422

    no_evidence = client.post(
        f"/jobs/{ids['second']}/mark-applied",
        json={"resume_used": "resumes/fullstack_resume.pdf", "application_method": "MANUAL"},
    )
    assert no_evidence.status_code == 422


def test_stats(api: tuple[TestClient, JobRepository, dict[str, int]]) -> None:
    client, _, _ = api
    response = client.get("/stats")

    assert response.status_code == 200
    assert response.json() == {
        "total_jobs": 2,
        "qualified": 1,
        "applied": 0,
        "needs_review": 0,
    }


def test_needs_review_is_counted_and_filterable(
    api: tuple[TestClient, JobRepository, dict[str, int]],
) -> None:
    client, repository, ids = api
    repository.mark_application_needs_review(ids["first"], "Complete platform login")

    stats = client.get("/stats")
    filtered = client.get("/jobs", params={"needs_review": "true"})

    assert stats.status_code == 200
    assert stats.json()["needs_review"] == 1
    assert stats.json()["qualified"] == 0
    assert filtered.status_code == 200
    assert [job["id"] for job in filtered.json()] == [ids["first"]]
    assert filtered.json()[0]["application_status"] == "NEEDS_REVIEW"
    assert filtered.json()[0]["review_reason"] == "Complete platform login"


def test_scrape_endpoint_uses_discovery_service(
    api: tuple[TestClient, JobRepository, dict[str, int]],
) -> None:
    client, repository, _ = api

    class FakeDiscoveryService:
        def __init__(self, received_repository: JobRepository, profile: UserProfile) -> None:
            assert received_repository is repository
            assert "Python" in profile.skills

        async def run(self, sources: object) -> DiscoverySummary:
            assert sources == ["configured-source"]
            return DiscoverySummary(1, 4, 2, 2, 0)

    client.app.state.sources_loader = lambda: ["configured-source"]  # type: ignore[attr-defined]
    client.app.state.discovery_service_factory = FakeDiscoveryService  # type: ignore[attr-defined]
    response = client.post("/scrape")

    assert response.status_code == 200
    assert response.json() == {
        "sources_checked": 1,
        "jobs_discovered": 4,
        "new_jobs": 2,
        "duplicates": 2,
        "failed_sources": 0,
        "remote_eligible": 0,
        "bengaluru_hybrid": 0,
        "bengaluru_onsite": 0,
        "filtered": 0,
        "jobs_scored": 1,
        "qualified": 1,
    }


def test_score_endpoint_updates_database(api: tuple[TestClient, JobRepository, dict[str, int]]) -> None:
    client, repository, ids = api
    response = client.post("/score")

    assert response.status_code == 200
    assert response.json()["jobs_scored"] == 1
    assert repository.get_job(ids["first"]).match_score > 70  # type: ignore[union-attr]
    assert repository.get_job(ids["first"]).status is JobStatus.QUALIFIED  # type: ignore[union-attr]


def test_cors_allows_only_local_frontend(
    api: tuple[TestClient, JobRepository, dict[str, int]],
) -> None:
    client, _, _ = api
    allowed = client.options(
        "/jobs",
        headers={
            "Origin": "http://localhost:5173",
            "Access-Control-Request-Method": "GET",
        },
    )
    blocked = client.options(
        "/jobs",
        headers={
            "Origin": "https://untrusted.example",
            "Access-Control-Request-Method": "GET",
        },
    )

    assert allowed.status_code == 200
    assert allowed.headers["access-control-allow-origin"] == "http://localhost:5173"
    assert blocked.status_code == 400


def test_application_endpoints_use_shared_service(
    api: tuple[TestClient, JobRepository, dict[str, int]],
) -> None:
    client, repository, ids = api

    class FakeApplicationService:
        def __init__(self) -> None:
            self.received: list[tuple[int | None, list[int] | None]] = []

        async def run_application_cycle(
            self, *, job_id: int | None = None, job_ids: list[int] | None = None
        ) -> ApplicationCycleSummary:
            self.received.append((job_id, job_ids))
            now = datetime.now(timezone.utc)
            return ApplicationCycleSummary(
                eligible=1,
                attempted=1,
                dry_run=True,
                started_at=now,
                completed_at=now,
            )

        def status(self) -> dict[str, object]:
            return {
                "running": False,
                "dry_run": True,
                "auto_apply_enabled": True,
                "last_run": None,
            }

    fake = FakeApplicationService()
    client.app.state.application_service = fake  # type: ignore[attr-defined]

    batch = client.post("/applications/run", json={"job_ids": [ids["first"]]})
    single = client.post(f"/applications/{ids['first']}/apply")
    status = client.get("/applications/status")

    assert batch.status_code == 200
    assert batch.json()["attempted"] == 1
    assert single.status_code == 200
    assert fake.received == [(None, [ids["first"]]), (ids["first"], None)]
    assert status.json()["dry_run"] is True
    assert repository.count_jobs() == 2
    assert client.post("/applications/999999/apply").status_code == 404


def test_connection_endpoints_use_session_manager(
    api: tuple[TestClient, JobRepository, dict[str, int]],
) -> None:
    client, _, _ = api

    class FakeSessionManager:
        def list_connections(self) -> list[ConnectionState]:
            return [ConnectionState("linkedin", "disconnected", False, "Connect")]

        def connect(self, platform: str) -> ConnectionState:
            assert platform == "linkedin"
            return ConnectionState(platform, "connecting", False, "Finish login")

        def clear_session(self, platform: str) -> None:
            assert platform == "linkedin"

    client.app.state.session_manager = FakeSessionManager()  # type: ignore[attr-defined]

    listed = client.get("/connections")
    started = client.post("/connections/linkedin/connect")
    cleared = client.delete("/connections/linkedin")

    assert listed.status_code == 200
    assert listed.json()[0]["connected"] is False
    assert started.status_code == 200
    assert started.json()["status"] == "connecting"
    assert cleared.status_code == 204


def test_resume_upload_parses_and_activates_docx(
    api: tuple[TestClient, JobRepository, dict[str, int]],
) -> None:
    client, repository, _ = api
    document = Document()
    document.add_paragraph("Technical Skills")
    document.add_paragraph("React, TypeScript")
    document.add_paragraph("Work Experience")
    document.add_paragraph("Frontend Developer")
    document.add_paragraph("5 years of experience in healthcare software")
    stream = BytesIO()
    document.save(stream)

    response = client.post(
        "/resume/upload",
        files={
            "file": (
                "rakshitha_resume.docx",
                stream.getvalue(),
                "application/vnd.openxmlformats-officedocument.wordprocessingml.document",
            )
        },
    )

    assert response.status_code == 200
    assert response.json()["filename"] == "rakshitha_resume.docx"
    assert response.json()["status"] == "Active"
    assert "React" in response.json()["skills"]
    assert client.get("/resume").json()["filename"] == "rakshitha_resume.docx"
    assert repository.get_active_resume().original_filename == "rakshitha_resume.docx"  # type: ignore[union-attr]
