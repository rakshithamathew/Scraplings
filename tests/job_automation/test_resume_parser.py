from __future__ import annotations

from pathlib import Path

from docx import Document

from job_automation.resume import extract_resume_text, parse_resume


def test_parse_docx_resume_without_inventing_missing_sections(tmp_path: Path) -> None:
    path = tmp_path / "candidate.docx"
    document = Document()
    document.add_paragraph("Frontend Developer")
    document.add_paragraph("5 years of experience building healthcare applications.")
    document.add_paragraph("Technical Skills")
    document.add_paragraph("React, TypeScript, Angular, FHIR")
    document.add_paragraph("Work Experience")
    document.add_paragraph("Senior Frontend Developer | Example Health")
    document.add_paragraph("Projects")
    document.add_paragraph("Clinical dashboard using React and TypeScript")
    document.add_paragraph("Education")
    document.add_paragraph("Bachelor of Engineering")
    document.save(path)

    profile = parse_resume(path)

    assert profile.years_of_experience == 5
    assert "React" in profile.technologies
    assert "FHIR" in profile.skills
    assert profile.job_titles == ["Senior Frontend Developer | Example Health"]
    assert profile.education == ["Bachelor of Engineering"]
    assert profile.certifications == []
    assert "healthcare" in profile.domain_experience
    assert profile.to_user_profile().minimum_experience == 5


def test_resume_parser_rejects_unsupported_files(tmp_path: Path) -> None:
    path = tmp_path / "resume.txt"
    path.write_text("Frontend Developer", encoding="utf-8")

    try:
        extract_resume_text(path)
    except ValueError as error:
        assert "PDF or DOCX" in str(error)
    else:
        raise AssertionError("Unsupported resume type was accepted")

