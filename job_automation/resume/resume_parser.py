"""Deterministic PDF/DOCX resume parsing without fabricated candidate facts."""

from __future__ import annotations

from collections import Counter
from pathlib import Path
import re

from pydantic import BaseModel, ConfigDict, Field

from job_automation.matching.ats_score import UserProfile


SUPPORTED_RESUME_SUFFIXES = {".pdf", ".docx"}
_SPACE = re.compile(r"\s+")
_YEAR_PATTERN = re.compile(r"\b(\d{1,2}(?:\.\d+)?)\s*\+?\s*years?\s+(?:of\s+)?(?:professional\s+|relevant\s+|industry\s+)?experience\b", re.I)
_TITLE_PATTERN = re.compile(
    r"\b(?:developer|engineer|architect|technical lead|team lead|manager|analyst|designer|consultant)\b",
    re.I,
)
_SECTION_NAMES = {
    "work_experience": {"experience", "work experience", "professional experience", "employment history"},
    "projects": {"projects", "project experience", "key projects"},
    "education": {"education", "academic background", "academics"},
    "certifications": {"certifications", "certificates", "licenses and certifications"},
    "skills": {"skills", "technical skills", "core competencies", "technologies", "tech stack"},
}
_TECHNOLOGIES = (
    "React", "React.js", "Angular", "Vue", "JavaScript", "TypeScript", "HTML", "CSS",
    "Redux", "Node.js", "Express", "Python", "Java", "C#", ".NET", "SQL", "GraphQL",
    "REST", "Azure", "AWS", "GCP", "Docker", "Kubernetes", "Git", "GitHub", "Cypress",
    "Playwright", "Jest", "Tailwind CSS", "Material UI", "Snowflake", "FHIR", "HL7",
    "OpenAI", "LLM", "RAG", "NLP", "OCR",
)
_DOMAINS = (
    "healthcare", "healthtech", "clinical", "medical", "fintech", "banking", "insurance",
    "e-commerce", "retail", "education", "telecom", "SaaS", "artificial intelligence",
    "machine learning", "generative AI", "GenAI",
)
_STOPWORDS = {
    "and", "the", "with", "for", "from", "that", "this", "have", "has", "using", "into",
    "your", "you", "our", "are", "was", "were", "will", "work", "experience", "years",
    "skills", "project", "projects", "education", "resume", "curriculum", "vitae",
}
_TITLE_LOCATION_SUFFIX = re.compile(
    r"(?:[,|\-]\s*)?(?:bengaluru|bangalore)(?:,?\s+(?:karnataka|india))?\s*$",
    re.I,
)


class ParsedCandidateProfile(BaseModel):
    """Facts extracted directly from one uploaded resume."""

    model_config = ConfigDict(extra="forbid")

    job_titles: list[str] = Field(default_factory=list)
    years_of_experience: float | None = None
    skills: list[str] = Field(default_factory=list)
    technologies: list[str] = Field(default_factory=list)
    work_experience: list[str] = Field(default_factory=list)
    projects: list[str] = Field(default_factory=list)
    education: list[str] = Field(default_factory=list)
    certifications: list[str] = Field(default_factory=list)
    domain_experience: list[str] = Field(default_factory=list)
    keywords: list[str] = Field(default_factory=list)
    resume_text: str = ""

    def to_user_profile(self) -> UserProfile:
        """Adapt extracted facts to the existing deterministic ATS scorer."""
        combined_skills = list(dict.fromkeys((*self.skills, *self.technologies)))
        years = int(self.years_of_experience) if self.years_of_experience is not None else None
        target_titles = [
            cleaned
            for title in self.job_titles
            if (cleaned := _TITLE_LOCATION_SUFFIX.sub("", title).strip(" ,-"))
        ]
        normalized_technologies = {value.casefold() for value in self.technologies}
        normalized_titles = " ".join(target_titles).casefold()
        related_titles: list[str] = []
        if {"react", "react.js", "angular", "javascript"} & normalized_technologies:
            related_titles.extend(
                ["Frontend Developer", "Frontend Engineer", "React Developer", "Angular Developer"]
            )
        if "full stack" in normalized_titles or {"node.js", "express"} & normalized_technologies:
            related_titles.extend(["Full Stack Developer", "Software Developer", "Software Engineer"])
        target_titles = list(dict.fromkeys((*target_titles, *related_titles)))
        return UserProfile(
            target_titles=target_titles,
            skills=combined_skills,
            preferred_skills=[],
            frontend_fullstack_skills=combined_skills,
            domain_keywords=self.domain_experience,
            preferred_locations=["Remote", "India", "Bengaluru", "Bangalore"],
            minimum_experience=years,
            maximum_experience=None,
            keywords=self.keywords,
            excluded_keywords=[],
        )


def _clean_line(value: str) -> str:
    return _SPACE.sub(" ", value).strip(" \t\r\n|•·-–—")


def extract_resume_text(path: str | Path) -> str:
    """Extract text from a supported resume, failing clearly for unreadable files."""
    resume_path = Path(path)
    suffix = resume_path.suffix.casefold()
    if suffix not in SUPPORTED_RESUME_SUFFIXES:
        raise ValueError("Resume must be a PDF or DOCX file")
    if not resume_path.is_file():
        raise ValueError(f"Resume file does not exist: {resume_path}")
    try:
        if suffix == ".pdf":
            from pypdf import PdfReader

            reader = PdfReader(str(resume_path))
            text = "\n".join(page.extract_text() or "" for page in reader.pages)
        else:
            from docx import Document

            document = Document(str(resume_path))
            paragraphs = [paragraph.text for paragraph in document.paragraphs]
            table_cells = [cell.text for table in document.tables for row in table.rows for cell in row.cells]
            text = "\n".join((*paragraphs, *table_cells))
    except ImportError as error:
        package = "pypdf" if suffix == ".pdf" else "python-docx"
        raise RuntimeError(f"{package} is required to parse {suffix} resumes") from error
    except Exception as error:
        raise ValueError(f"Unable to read resume: {error}") from error
    cleaned = "\n".join(line for raw in text.splitlines() if (line := _clean_line(raw)))
    if not cleaned:
        raise ValueError("No readable text was found in the resume")
    return cleaned


def _sections(lines: list[str]) -> dict[str, list[str]]:
    result = {name: [] for name in _SECTION_NAMES}
    current: str | None = None
    for line in lines:
        heading = line.casefold().rstrip(":")
        matched = next((name for name, names in _SECTION_NAMES.items() if heading in names), None)
        if matched:
            current = matched
        elif current:
            result[current].append(line)
    return result


def _unique_present(terms: tuple[str, ...], text: str) -> list[str]:
    found: list[str] = []
    for term in terms:
        if re.search(r"(?<!\w)" + re.escape(term) + r"(?!\w)", text, re.I):
            found.append(term)
    return found


def parse_resume(path: str | Path) -> ParsedCandidateProfile:
    """Parse one resume into a normalized, cacheable candidate profile."""
    text = extract_resume_text(path)
    lines = text.splitlines()
    sections = _sections(lines)
    technologies = _unique_present(_TECHNOLOGIES, text)
    domains = _unique_present(_DOMAINS, text)

    explicit_skills: list[str] = []
    for line in sections["skills"]:
        for item in re.split(r"[,;|•]+", line):
            cleaned = _clean_line(item)
            if cleaned and len(cleaned) <= 60 and len(cleaned.split()) <= 6:
                explicit_skills.append(cleaned)
    skills = list(dict.fromkeys((*explicit_skills, *technologies)))

    titles: list[str] = []
    title_source = sections["work_experience"] or lines
    for line in title_source:
        if _TITLE_PATTERN.search(line) and len(line.split()) <= 12:
            titles.append(line)
    titles = list(dict.fromkeys(titles))[:20]

    years = [float(match) for match in _YEAR_PATTERN.findall(text)]
    tokens = re.findall(r"[A-Za-z][A-Za-z+#.]{2,}", text)
    keyword_counts = Counter(token.casefold() for token in tokens if token.casefold() not in _STOPWORDS)
    keywords = [token for token, _ in keyword_counts.most_common(40)]

    return ParsedCandidateProfile(
        job_titles=titles,
        years_of_experience=max(years) if years else None,
        skills=skills,
        technologies=technologies,
        work_experience=sections["work_experience"],
        projects=sections["projects"],
        education=sections["education"],
        certifications=sections["certifications"],
        domain_experience=domains,
        keywords=keywords,
        resume_text=text,
    )
